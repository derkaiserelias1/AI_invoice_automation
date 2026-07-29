"""Unified text extraction for invoices: PDF, images (EasyOCR), Office, plain text."""

from __future__ import annotations

import csv
import io
from pathlib import Path

# extension -> short file_type label stored on Document
ALLOWED_EXTENSIONS: dict[str, str] = {
    ".pdf": "pdf",
    ".png": "png",
    ".jpg": "jpg",
    ".jpeg": "jpg",
    ".webp": "webp",
    ".tif": "tiff",
    ".tiff": "tiff",
    ".bmp": "bmp",
    ".docx": "docx",
    ".txt": "txt",
    ".csv": "csv",
    ".xlsx": "xlsx",
}

IMAGE_TYPES = {"png", "jpg", "webp", "tiff", "bmp"}

# If embedded PDF text is shorter than this, run OCR on page images
PDF_TEXT_MIN_CHARS = 40

_ocr_reader = None


def allowed_extensions_accept() -> str:
    """HTML accept= attribute value."""
    parts = sorted({f"application/pdf", "text/plain", "text/csv", "image/*"})
    exts = ",".join(sorted(ALLOWED_EXTENSIONS.keys()))
    return f"{exts},.jpeg"


def file_type_from_name(filename: str) -> str | None:
    ext = Path(filename).suffix.lower()
    return ALLOWED_EXTENSIONS.get(ext)


def is_allowed_filename(filename: str) -> bool:
    return file_type_from_name(filename) is not None


def _get_ocr():
    """Lazy-load EasyOCR (heavy models; first call may download)."""
    global _ocr_reader
    if _ocr_reader is None:
        import easyocr

        # English + common latin; add more langs later if needed
        _ocr_reader = easyocr.Reader(["en"], gpu=False, verbose=False)
    return _ocr_reader


def ocr_image_bytes(image_bytes: bytes) -> str:
    import numpy as np
    from PIL import Image

    img = Image.open(io.BytesIO(image_bytes)).convert("RGB")
    arr = np.array(img)
    reader = _get_ocr()
    lines = reader.readtext(arr, detail=0, paragraph=True)
    if isinstance(lines, list):
        return "\n".join(str(x) for x in lines if x).strip()
    return str(lines or "").strip()


def ocr_image_path(path: Path) -> str:
    return ocr_image_bytes(path.read_bytes())


def _extract_pdf_embedded(path: Path) -> str:
    import pdfplumber

    parts: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                parts.append(t)
    return "\n".join(parts).strip()


def _extract_pdf_ocr(path: Path, max_pages: int = 10) -> str:
    """Render PDF pages with PyMuPDF and OCR each page."""
    import fitz  # PyMuPDF

    doc = fitz.open(path)
    parts: list[str] = []
    try:
        n = min(len(doc), max_pages)
        for i in range(n):
            page = doc[i]
            # ~150 dpi-ish for speed/accuracy balance
            mat = fitz.Matrix(2.0, 2.0)
            pix = page.get_pixmap(matrix=mat, alpha=False)
            text = ocr_image_bytes(pix.tobytes("png"))
            if text:
                parts.append(text)
    finally:
        doc.close()
    return "\n\n".join(parts).strip()


def extract_pdf(path: Path) -> str:
    embedded = _extract_pdf_embedded(path)
    if len(embedded) >= PDF_TEXT_MIN_CHARS:
        return embedded
    # Scanned / empty text layer
    ocr_text = _extract_pdf_ocr(path)
    if ocr_text:
        if embedded:
            return f"{embedded}\n\n{ocr_text}".strip()
        return ocr_text
    return embedded


def extract_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    # tables
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def extract_txt(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc).strip()
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace").strip()


def extract_csv(path: Path) -> str:
    text = extract_txt(path)
    # Normalize to a stable line layout for the LLM
    try:
        reader = csv.reader(io.StringIO(text))
        rows = [", ".join(cell.strip() for cell in row if cell is not None) for row in reader]
        return "\n".join(r for r in rows if r.strip()).strip() or text
    except Exception:
        return text


def extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for sheet in wb.worksheets:
            parts.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    parts.append(" | ".join(cells))
    finally:
        wb.close()
    return "\n".join(parts).strip()


def extract_text_from_file(path: str | Path, file_type: str | None = None) -> str:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(str(path))

    ft = (file_type or file_type_from_name(path.name) or path.suffix.lstrip(".")).lower()
    if ft == "jpeg":
        ft = "jpg"

    if ft == "pdf":
        return extract_pdf(path)
    if ft in IMAGE_TYPES:
        return ocr_image_path(path)
    if ft == "docx":
        return extract_docx(path)
    if ft == "txt":
        return extract_txt(path)
    if ft == "csv":
        return extract_csv(path)
    if ft == "xlsx":
        return extract_xlsx(path)

    raise ValueError(f"Unsupported file type: {ft}")
